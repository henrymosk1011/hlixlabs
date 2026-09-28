#!/usr/bin/env python3
"""
Exports data/products.json to an .xlsx for the site owner to edit and hand
back as the new base (Powder-form) price list. Every variant (dose size) is
its own row, matching what a shopper actually picks on the site.

Run: ./.venv/bin/python scripts/export_pricelist.py
Writes: data/price_export.xlsx
"""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
PRODUCTS = ROOT / "data" / "products.json"
OUT = ROOT / "data" / "price_export.xlsx"


def main():
    products = json.loads(PRODUCTS.read_text())
    products = sorted(products, key=lambda p: (p["category_label"], p["name"], float(p["dose_amount"] or 0)))

    wb = Workbook()
    ws = wb.active
    ws.title = "Prices"

    headers = ["Category", "Product name", "Dose", "SKU", "Current price (Powder base)"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for p in products:
        ws.append([
            p["category_label"],
            p["name"],
            p["dose"],
            p["sku"],
            float(p["price"]),
        ])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = "$#,##0.00"

    widths = [24, 34, 10, 16, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:E{ws.max_row}"

    wb.save(OUT)
    print(f"Wrote {len(products)} rows to {OUT}")


if __name__ == "__main__":
    main()
